"""Export saved workbook content into a deployable, validated JSON overlay.

The source workbook and the catalog's winners/participants are never modified.
Run with --dry-run to inspect the match report without writing anything.
"""
import argparse
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import unicodedata
from urllib.parse import urlsplit

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / 'data' / 'workbook_content.json'
FIELDS = {
    'week number': 'week', 'competition type': 'type', 'competition name': 'name',
    'competition family': 'family', 'competition description': 'description',
    'competition strategy': 'strategy', 'photo of competition': 'image_url',
    'video of competition': 'video_url', 'guest appearances': 'guest_appearances',
}
# An existing catalog spelling for the same named recurring format.
FAMILY_ALIASES = {'sequences & explosions': 'Sequence Explosion'}


def normalize(value):
    return ' '.join(unicodedata.normalize('NFKC', str(value or '')).casefold().split())


def text_value(cell):
    if cell.data_type in ('f', 'e'):
        raise ValueError(f'{cell.parent.title}!{cell.coordinate}: formulas/errors require review')
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def web_url(value, location):
    parsed = urlsplit(value)
    if value and (parsed.scheme not in ('http', 'https') or not parsed.netloc
                  or parsed.username or parsed.password):
        raise ValueError(f'{location}: expected a public http(s) link, got {value!r}')
    return value


def source_urls(cell):
    if not cell.comment:
        return []
    urls = re.findall(r'https?://[^\s<>|]+', cell.comment.text)
    return sorted({web_url(url.rstrip('.,;'), cell.coordinate) for url in urls})


def merge_nonblank(previous, current):
    result = dict(previous)
    for key, value in current.items():
        if value not in ('', None, []):
            result[key] = value
    return result


def build_content(workbook_path, catalog, previous=None):
    previous = previous or {}
    raw = Path(workbook_path).read_bytes()
    book = load_workbook(BytesIO(raw), data_only=False)
    exact = defaultdict(list)
    for event in catalog['appearances']:
        key = (event['season'], normalize(event['week']), normalize(event['type']), normalize(event['variation']))
        exact[key].append(event)
    known_names = {normalize(c['name']): c['name'] for c in catalog['competitions']}

    def family_name(name):
        alias = FAMILY_ALIASES.get(normalize(name), name)
        return known_names.get(normalize(alias), alias)

    def family_list(value):
        if not value or normalize(value) in ('none', 'n/a'):
            return []
        value = family_name(value)
        if normalize(value) in known_names:
            return [family_name(value)]
        # Some family names themselves contain commas ("Ready, Set, Woah").
        # Split only when the entire value can be segmented into known names.
        def segment(remaining):
            if normalize(remaining) in known_names:
                return [family_name(remaining)]
            for name in sorted(known_names.values(), key=len, reverse=True):
                prefix = name + ', '
                if normalize(remaining).startswith(normalize(prefix) + ' '):
                    tail = segment(remaining[len(prefix):])
                    if tail:
                        return [name] + tail
            return None
        return segment(value) or [value]

    families = dict(previous.get('families', {}))
    family_sheet = book['Family Database']
    family_headers = {normalize(c.value): c.column for c in family_sheet[1] if c.value}
    if not {'family name', 'description', 'strategy'} <= family_headers.keys():
        raise ValueError('Family Database must have Family Name, Description, and Strategy headers')
    imported_families = []
    for row in range(2, family_sheet.max_row + 1):
        values = {key: text_value(family_sheet.cell(row, col)) for key, col in family_headers.items()}
        if not values['family name']:
            continue
        name = family_name(values['family name'])
        entry = {'name': name, 'workbook_name': values['family name'], 'description': values['description'],
                 'strategy': values['strategy']}
        families[name] = merge_nonblank(families.get(name, {}), entry)
        imported_families.append(name)
        known_names[normalize(name)] = name

    events = dict(previous.get('events', {}))
    seen = {}
    report = {'rows': 0, 'families': len(imported_families), 'seasons': {}, 'resolved_multi_matches': [], 'errors': []}
    for sheet in book:
        match = re.fullmatch(r'Season (\d+)', sheet.title)
        if not match:
            continue
        season = int(match[1])
        if season < 2:
            raise ValueError('Season 1 is intentionally excluded')
        headers = {normalize(c.value).replace(' (links)', ''): c.column for c in sheet[1] if c.value}
        if not set(list(FIELDS)[:6]) <= headers.keys():
            raise ValueError(f'{sheet.title}: missing required competition headers')
        season_count = 0
        for row in range(2, sheet.max_row + 1):
            values = {}
            sources = {}
            for header, field in FIELDS.items():
                if header not in headers:
                    continue
                cell = sheet.cell(row, headers[header])
                value = text_value(cell)
                if field in ('image_url', 'video_url'):
                    value = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else value
                    value = web_url(value, f'{sheet.title}!{cell.coordinate}')
                values[field] = value
                urls = source_urls(cell)
                if urls:
                    sources[field] = urls
            if not any(values.values()):
                continue
            location = f'{sheet.title}!{row}'
            if not all(values.get(key) for key in ('name', 'week', 'type')):
                report['errors'].append(f'{location}: missing name, week, or competition type')
                continue
            candidates = exact.get((season, normalize(values['week']), normalize(values['type']), normalize(values['name'])), [])
            event_keys = {c['event_key'] for c in candidates}
            resolved_family = family_name(values.get('family', ''))
            resolved_families = family_list(resolved_family)
            if len(event_keys) > 1:
                family_candidates = [c for c in candidates if normalize(c['competition']) in {normalize(f) for f in resolved_families}]
                if len({c['event_key'] for c in family_candidates}) == 1:
                    candidates = family_candidates
                    event_keys = {c['event_key'] for c in candidates}
                    report['resolved_multi_matches'].append(location)
            if len(event_keys) != 1:
                report['errors'].append(f'{location}: {len(event_keys)} matching events for {values["name"]!r}; requires review')
                continue
            event_key = next(iter(event_keys))
            if event_key in seen:
                report['errors'].append(f'{location}: duplicates {seen[event_key]} ({event_key})')
                continue
            seen[event_key] = location
            values.update({'family': resolved_family, 'families': resolved_families, 'season': season, 'sheet': sheet.title, 'row': row})
            old = events.get(event_key, {})
            values['sources'] = {**old.get('sources', {}), **sources}
            events[event_key] = merge_nonblank(old, values)
            report['rows'] += 1
            season_count += 1
        report['seasons'][str(season)] = season_count
    book.close()
    if report['errors']:
        raise ValueError('\n'.join(report['errors']))
    if not report['rows']:
        raise ValueError('No competition rows found')
    return {
        'version': 1, 'workbook_name': Path(workbook_path).name,
        'workbook_sha256': hashlib.sha256(raw).hexdigest(),
        'imported_at': datetime.now(timezone.utc).isoformat(),
        'families': families, 'events': events,
    }, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--catalog', type=Path, default=ROOT / 'data/us_recurring_competitions.json')
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()
    previous = json.loads(args.output.read_text()) if args.output.exists() else None
    try:
        content, report = build_content(args.workbook, json.loads(args.catalog.read_text()), previous)
    except ValueError as error:
        parser.exit(1, f'Import aborted; no files changed:\n{error}\n')
    print(json.dumps(report, indent=2))
    if not args.dry_run:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        temporary = args.output.with_suffix('.json.tmp')
        temporary.write_text(json.dumps(content, ensure_ascii=False, indent=2) + '\n')
        temporary.replace(args.output)
        print(f'Wrote {args.output}')


if __name__ == '__main__':
    main()
